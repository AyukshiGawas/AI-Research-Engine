"""DOCX text extractor implementation using python-docx."""

from pathlib import Path
import docx

from app.exceptions import DocumentExtractionError
from app.services.extractors.base import BaseExtractor, ExtractionResult


class DOCXExtractor(BaseExtractor):
    """Text extractor for Office Open XML Document (.docx) files using python-docx."""

    def extract(self, file_path: Path) -> ExtractionResult:
        try:
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            full_text = "\n\n".join(paragraphs).strip()
            word_count = len(full_text.split()) if full_text else 0
            return ExtractionResult(
                text=full_text,
                page_count=None,
                word_count=word_count,
            )
        except Exception as exc:
            raise DocumentExtractionError(f"Failed to extract text from DOCX file: {exc}") from exc
